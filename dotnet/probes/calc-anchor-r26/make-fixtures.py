"""Builds the three sheet fixtures for the 26.8.1.1 anchor / print-range round.

sheet-chart-editas-onecell.xlsx: one bar chart whose twoCellAnchor editAs="oneCell" spans
  columns B..J (from/to ~ 9 columns) while its xdr:xfrm/a:ext says 3 x 2 inches. 26.8.1.1 keeps
  the stated size (drawingfragment.cxx:284-295).
The other two are copies of the analysis probes (rot90-oneCell.xlsx, hidden-printarea.xlsx).
Run with python3 -I make-fixtures.py <outdir>.
"""
import io, re, shutil, sys, zipfile
import openpyxl
from openpyxl.chart import BarChart, Reference

out = sys.argv[1]
wb = openpyxl.Workbook(); ws = wb.active
for i in range(1, 6):
    ws.cell(row=i, column=1, value=i * 10)
chart = BarChart(); chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=5))
ws.add_chart(chart, "B2")
buf = io.BytesIO(); wb.save(buf)

EMU = 914400
anchor = ('<xdr:twoCellAnchor editAs="oneCell" '
          'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
          '<xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
          '<xdr:to><xdr:col>10</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>30</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
          '<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="2" name="Chart 1"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>'
          f'<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="{3 * EMU}" cy="{2 * EMU}"/></xdr:xfrm>'
          '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
          '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
          '</a:graphicData></a:graphic></xdr:graphicFrame><xdr:clientData/></xdr:twoCellAnchor>')

src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
with zipfile.ZipFile(f"{out}/sheet-chart-editas-onecell.xlsx", "w", zipfile.ZIP_DEFLATED) as dst:
    for name in src.namelist():
        data = src.read(name)
        if name == "xl/drawings/drawing1.xml":
            text = data.decode()
            text = re.sub(r"<(xdr:)?oneCellAnchor>.*</(xdr:)?oneCellAnchor>", anchor, text, flags=re.S)
            assert "editAs" in text, text
            data = text.encode()
        dst.writestr(name, data)

for name, target in (("rot90-oneCell.xlsx", "sheet-drawing-rot90-onecell.xlsx"),
                     ("hidden-printarea.xlsx", "sheet-hidden-print-area.xlsx")):
    shutil.copy(f"/home/user/lo-upstream/analysis/calc-probes/{name}", f"{out}/{target}")
